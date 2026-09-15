import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from paper_dataset import papers

def create_presentation():
    prs = Presentation()
    
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Color Palette
    c_navy = RGBColor(15, 23, 42)      # #0f172a
    c_slate = RGBColor(30, 41, 59)     # #1e293b
    c_blue = RGBColor(2, 132, 199)     # #0284c7
    c_emerald = RGBColor(5, 150, 105)  # #059669
    c_light_bg = RGBColor(248, 250, 252) # #f8fafc
    c_white = RGBColor(255, 255, 255)
    c_gray_text = RGBColor(100, 116, 139) # #64748b
    c_border = RGBColor(226, 232, 240)

    blank_layout = prs.slide_layouts[6]

    # Helper: Add Slide Header
    def add_slide_header(slide, title_text, category_tag="Academic Survey"):
        # Header background banner
        header_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        header_rect.fill.solid()
        header_rect.fill.fore_color.rgb = c_navy
        header_rect.line.fill.background()

        # Title text
        txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(10), Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = c_white
        p.font.name = "Arial"

        # Tag
        txBox2 = slide.shapes.add_textbox(Inches(10.5), Inches(0.3), Inches(2.3), Inches(0.5))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = f"HomeWise AI | {category_tag}"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = c_emerald
        p2.alignment = PP_ALIGN.RIGHT

    # Helper: Format Table Cell Text
    def format_cell(cell, text, bold=False, color=c_slate, size=10, align=PP_ALIGN.LEFT, bg_color=None):
        if bg_color:
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(0.06)
        tf.margin_bottom = Inches(0.06)
        
        p = tf.paragraphs[0]
        p.text = str(text)
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_navy
    bg1.line.fill.background()

    # Decorative Accent Bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(0.15), Inches(3.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = c_blue
    bar.line.fill.background()

    # Main Title & Subtitle Box
    tbox1 = slide1.shapes.add_textbox(Inches(1.6), Inches(1.8), Inches(10.5), Inches(3.8))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "Literature Review"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = c_white
    p1.font.name = "Arial"
    p1.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "Analysis of 15 Verified Academic Research Papers"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = c_emerald
    p2.font.name = "Arial"
    p2.space_after = Pt(18)

    p3 = tf1.add_paragraph()
    p3.text = "Project Title: HomeWise AI – AI-Based Neighborhood Intelligence for Home-Buying Decision Support\nAcademic Review 2 Evaluation | Comprehensive Literature Survey & Research Gap Synthesis"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.font.name = "Arial"

    # =========================================================================
    # SLIDE 2: PAPERS 1 - 5
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide2, "Literature Review – Papers 1–5", "Part 1 of 3")

    rows = 6
    cols = 7
    left = Inches(0.5)
    top = Inches(1.3)
    width = Inches(12.333)
    height = Inches(5.7)

    table_shape2 = slide2.shapes.add_table(rows, cols, left, top, width, height)
    t2 = table_shape2.table
    
    col_widths2 = [Inches(0.5), Inches(2.4), Inches(1.6), Inches(1.8), Inches(0.7), Inches(2.6), Inches(2.733)]
    for idx, w in enumerate(col_widths2):
        t2.columns[idx].width = w

    headers = ["No.", "Paper Title", "Authors", "Journal / Venue", "Year", "Methodology", "Relevance & Limitations"]
    for idx, h in enumerate(headers):
        format_cell(t2.cell(0, idx), h, bold=True, color=c_white, size=11, align=PP_ALIGN.CENTER if idx in [0,4] else PP_ALIGN.LEFT, bg_color=c_navy)

    p_group1 = papers[0:5]
    for r_idx, p in enumerate(p_group1, 1):
        bg_row = c_light_bg if r_idx % 2 == 1 else c_white
        format_cell(t2.cell(r_idx, 0), str(p["no"]), bold=True, size=10, align=PP_ALIGN.CENTER, bg_color=bg_row)
        format_cell(t2.cell(r_idx, 1), p["title"], bold=True, size=10, bg_color=bg_row)
        format_cell(t2.cell(r_idx, 2), p["authors_short"], size=9, bg_color=bg_row)
        format_cell(t2.cell(r_idx, 3), p["journal"], size=9, bg_color=bg_row)
        format_cell(t2.cell(r_idx, 4), str(p["year"]), size=10, align=PP_ALIGN.CENTER, bg_color=bg_row)
        format_cell(t2.cell(r_idx, 5), p["method"], size=9, bg_color=bg_row)
        rel_text = f"Relevance: {p['relevance']}\nLimitation: {p['limitation']}"
        format_cell(t2.cell(r_idx, 6), rel_text, size=8, bg_color=bg_row)

    # =========================================================================
    # SLIDE 3: PAPERS 6 - 10
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide3, "Literature Review – Papers 6–10", "Part 2 of 3")

    table_shape3 = slide3.shapes.add_table(rows, cols, left, top, width, height)
    t3 = table_shape3.table
    for idx, w in enumerate(col_widths2):
        t3.columns[idx].width = w

    for idx, h in enumerate(headers):
        format_cell(t3.cell(0, idx), h, bold=True, color=c_white, size=11, align=PP_ALIGN.CENTER if idx in [0,4] else PP_ALIGN.LEFT, bg_color=c_navy)

    p_group2 = papers[5:10]
    for r_idx, p in enumerate(p_group2, 1):
        bg_row = c_light_bg if r_idx % 2 == 1 else c_white
        format_cell(t3.cell(r_idx, 0), str(p["no"]), bold=True, size=10, align=PP_ALIGN.CENTER, bg_color=bg_row)
        format_cell(t3.cell(r_idx, 1), p["title"], bold=True, size=10, bg_color=bg_row)
        format_cell(t3.cell(r_idx, 2), p["authors_short"], size=9, bg_color=bg_row)
        format_cell(t3.cell(r_idx, 3), p["journal"], size=9, bg_color=bg_row)
        format_cell(t3.cell(r_idx, 4), str(p["year"]), size=10, align=PP_ALIGN.CENTER, bg_color=bg_row)
        format_cell(t3.cell(r_idx, 5), p["method"], size=9, bg_color=bg_row)
        rel_text = f"Relevance: {p['relevance']}\nLimitation: {p['limitation']}"
        format_cell(t3.cell(r_idx, 6), rel_text, size=8, bg_color=bg_row)

    # =========================================================================
    # SLIDE 4: PAPERS 11 - 15
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide4, "Literature Review – Papers 11–15", "Part 3 of 3")

    table_shape4 = slide4.shapes.add_table(rows, cols, left, top, width, height)
    t4 = table_shape4.table
    for idx, w in enumerate(col_widths2):
        t4.columns[idx].width = w

    for idx, h in enumerate(headers):
        format_cell(t4.cell(0, idx), h, bold=True, color=c_white, size=11, align=PP_ALIGN.CENTER if idx in [0,4] else PP_ALIGN.LEFT, bg_color=c_navy)

    p_group3 = papers[10:15]
    for r_idx, p in enumerate(p_group3, 1):
        bg_row = c_light_bg if r_idx % 2 == 1 else c_white
        format_cell(t4.cell(r_idx, 0), str(p["no"]), bold=True, size=10, align=PP_ALIGN.CENTER, bg_color=bg_row)
        format_cell(t4.cell(r_idx, 1), p["title"], bold=True, size=10, bg_color=bg_row)
        format_cell(t4.cell(r_idx, 2), p["authors_short"], size=9, bg_color=bg_row)
        format_cell(t4.cell(r_idx, 3), p["journal"], size=9, bg_color=bg_row)
        format_cell(t4.cell(r_idx, 4), str(p["year"]), size=10, align=PP_ALIGN.CENTER, bg_color=bg_row)
        format_cell(t4.cell(r_idx, 5), p["method"], size=9, bg_color=bg_row)
        rel_text = f"Relevance: {p['relevance']}\nLimitation: {p['limitation']}"
        format_cell(t4.cell(r_idx, 6), rel_text, size=8, bg_color=bg_row)

    # =========================================================================
    # SLIDE 5: RESEARCH GAP & PROPOSED CONTRIBUTION
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide5, "Research Gap & Proposed Contribution", "Synthesis Architecture")

    # Flow Box 1: Existing Research
    b1 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.5), Inches(3.6), Inches(5.2))
    b1.fill.solid()
    b1.fill.fore_color.rgb = c_light_bg
    b1.line.color.rgb = c_border
    tf_b1 = b1.text_frame
    tf_b1.word_wrap = True
    tf_b1.margin_left = Inches(0.2)
    tf_b1.margin_right = Inches(0.2)
    
    p = tf_b1.paragraphs[0]
    p.text = "Existing Research Landscape"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_navy

    bullets_b1 = [
        "GIS Suitability Analysis (Malczewski; Chandio)",
        "MCDM Criteria Weighting (Zolfani; Pamučar)",
        "Machine Learning Valuation (Baldominos)",
        "Hedonic Environmental Pricing (Zhang & Dong)",
        "Transit Infrastructure Access (Foth et al.)",
        "Consumer Choice Models (Schirmer; Walker)"
    ]
    for b in bullets_b1:
        p_b = tf_b1.add_paragraph()
        p_b.text = f"• {b}"
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = c_slate
        p_b.space_before = Pt(8)

    # Arrow 1
    arr1 = slide5.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.35), Inches(3.8), Inches(0.6), Inches(0.4))
    arr1.fill.solid()
    arr1.fill.fore_color.rgb = c_blue
    arr1.line.fill.background()

    # Flow Box 2: Limitation & Research Gap
    b2 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.1), Inches(1.5), Inches(3.6), Inches(5.2))
    b2.fill.solid()
    b2.fill.fore_color.rgb = RGBColor(254, 243, 199) # soft amber
    b2.line.color.rgb = RGBColor(251, 191, 36)
    tf_b2 = b2.text_frame
    tf_b2.word_wrap = True
    tf_b2.margin_left = Inches(0.2)
    tf_b2.margin_right = Inches(0.2)

    p = tf_b2.paragraphs[0]
    p.text = "Identified Research Gap"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(146, 64, 14)

    bullets_b2 = [
        "Fragmented Silos: GIS, MCDM, and ML models are studied in isolation.",
        "Lack of Consumer Accessibility: Complex desktop GIS/MCDM software built for planners, not home buyers.",
        "Black-Box ML Predictions: ML models lack natural-language explanation rationale.",
        "Static Preference Models: Stated choice surveys do not support interactive dynamic re-weighting."
    ]
    for b in bullets_b2:
        p_b = tf_b2.add_paragraph()
        p_b.text = f"⚠ {b}"
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = RGBColor(120, 53, 15)
        p_b.space_before = Pt(8)

    # Arrow 2
    arr2 = slide5.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.85), Inches(3.8), Inches(0.6), Inches(0.4))
    arr2.fill.solid()
    arr2.fill.fore_color.rgb = c_emerald
    arr2.line.fill.background()

    # Flow Box 3: HomeWise AI Contribution
    b3 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.6), Inches(1.5), Inches(3.2), Inches(5.2))
    b3.fill.solid()
    b3.fill.fore_color.rgb = RGBColor(209, 250, 229) # soft emerald
    b3.line.color.rgb = RGBColor(52, 211, 153)
    tf_b3 = b3.text_frame
    tf_b3.word_wrap = True
    tf_b3.margin_left = Inches(0.2)
    tf_b3.margin_right = Inches(0.2)

    p = tf_b3.paragraphs[0]
    p.text = "HomeWise AI Integrated Framework"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(6, 78, 59)

    bullets_b3 = [
        "Unified Decision Pipeline: Combines buyer preferences + budget constraints + 9 neighborhood factors.",
        "Interactive Web GIS Map: Leaflet spatial visual map with 40 locality markers.",
        "Transparent MCDM Ranking: Dynamic weighted scoring engine.",
        "Prototype AI Insight: Generates natural language explanation text."
    ]
    for b in bullets_b3:
        p_b = tf_b3.add_paragraph()
        p_b.text = f"✓ {b}"
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = RGBColor(6, 95, 70)
        p_b.space_before = Pt(8)

    # =========================================================================
    # SLIDE 6: COMPACT SUMMARY (PAPERS 1 - 8)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide6, "Literature Review – Summary (Papers 1–8)", "Compact Overview")

    table_shape6 = slide6.shapes.add_table(9, 4, Inches(0.5), Inches(1.3), Inches(12.333), Inches(5.7))
    t6 = table_shape6.table
    col_w6 = [Inches(0.6), Inches(2.2), Inches(4.5), Inches(5.033)]
    for idx, w in enumerate(col_w6): t6.columns[idx].width = w

    headers_summary = ["No.", "Author & Year", "Paper Title", "Methodology / Main Contribution"]
    for idx, h in enumerate(headers_summary):
        format_cell(t6.cell(0, idx), h, bold=True, color=c_white, size=11, align=PP_ALIGN.CENTER if idx==0 else PP_ALIGN.LEFT, bg_color=c_navy)

    for r_idx, p in enumerate(papers[0:8], 1):
        bg_row = c_light_bg if r_idx % 2 == 1 else c_white
        format_cell(t6.cell(r_idx, 0), str(p["no"]), bold=True, size=10, align=PP_ALIGN.CENTER, bg_color=bg_row)
        format_cell(t6.cell(r_idx, 1), f"{p['authors_short']} ({p['year']})", bold=True, size=10, bg_color=bg_row)
        format_cell(t6.cell(r_idx, 2), p["title"], size=9.5, bg_color=bg_row)
        contrib_text = f"Method: {p['method_type']}. {p['findings']}"
        format_cell(t6.cell(r_idx, 3), contrib_text, size=9, bg_color=bg_row)

    # =========================================================================
    # SLIDE 7: COMPACT SUMMARY (PAPERS 9 - 15)
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide7, "Literature Review – Summary (Papers 9–15)", "Compact Overview")

    table_shape7 = slide7.shapes.add_table(8, 4, Inches(0.5), Inches(1.3), Inches(12.333), Inches(5.5))
    t7 = table_shape7.table
    for idx, w in enumerate(col_w6): t7.columns[idx].width = w

    for idx, h in enumerate(headers_summary):
        format_cell(t7.cell(0, idx), h, bold=True, color=c_white, size=11, align=PP_ALIGN.CENTER if idx==0 else PP_ALIGN.LEFT, bg_color=c_navy)

    for r_idx, p in enumerate(papers[8:15], 1):
        bg_row = c_light_bg if r_idx % 2 == 1 else c_white
        format_cell(t7.cell(r_idx, 0), str(p["no"]), bold=True, size=10, align=PP_ALIGN.CENTER, bg_color=bg_row)
        format_cell(t7.cell(r_idx, 1), f"{p['authors_short']} ({p['year']})", bold=True, size=10, bg_color=bg_row)
        format_cell(t7.cell(r_idx, 2), p["title"], size=9.5, bg_color=bg_row)
        contrib_text = f"Method: {p['method_type']}. {p['findings']}"
        format_cell(t7.cell(r_idx, 3), contrib_text, size=9, bg_color=bg_row)

    ppt_path = "HomeWise_AI_Literature_Review_15_Papers.pptx"
    prs.save(ppt_path)
    print(f"Presentation saved successfully to {ppt_path}")

if __name__ == "__main__":
    create_presentation()
