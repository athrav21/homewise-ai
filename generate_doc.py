import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from paper_dataset import papers

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins (1 inch everywhere)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Color Palette Constants
    c_navy_hex = "0F172A"
    c_slate_hex = "1E293B"
    c_blue_hex = "0284C7"
    c_emerald_hex = "059669"
    c_bg_hex = "F8FAFC"
    c_amber_bg = "FEF3C7"

    c_navy_rgb = RGBColor(15, 23, 42)
    c_blue_rgb = RGBColor(2, 132, 199)
    c_emerald_rgb = RGBColor(5, 150, 105)
    c_slate_rgb = RGBColor(30, 41, 59)

    # Styles setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = c_slate_rgb

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("HomeWise AI — Comprehensive Academic Literature Review")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = c_navy_rgb
    p_title.paragraph_format.space_after = Pt(4)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("AI-Based Neighborhood Intelligence for Home-Buying Decision Support\nReview 2 Academic Literature Survey of 15 Verified Publications")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = c_emerald_rgb
    r_sub.font.bold = True
    p_sub.paragraph_format.space_after = Pt(24)

    # Section 1: Introduction
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Introduction to the Literature Review")
    r_h1.font.color.rgb = c_navy_rgb
    r_h1.font.bold = True

    p_intro1 = doc.add_paragraph(
        "Residential location selection is a complex, high-involvement decision problem that sits at the intersection of urban economics, "
        "geospatial decision science, computer science, and consumer behavioral psychology. When selecting a home, buyers do not simply purchase "
        "a physical structure; they invest in a neighborhood environment defined by access to quality schools, emergency healthcare facilities, "
        "commercial markets, public transportation, air quality, flood safety, and long-term community security. Over the past three decades, "
        "academic literature has examined various dimensions of this process using discrete choice modeling, GIS-based multi-criteria evaluation, "
        "hedonic property pricing, and machine learning."
    )
    p_intro1.paragraph_format.space_after = Pt(8)
    p_intro1.paragraph_format.line_spacing = 1.15

    p_intro2 = doc.add_paragraph(
        "This literature survey systematically reviews 15 academically verified research papers directly relevant to the HomeWise AI decision-support "
        "system. The review evaluates established methodologies, identifies critical limitations in existing analytical models, and outlines the "
        "specific research gap that HomeWise AI addresses by integrating interactive GIS web mapping, dynamic multi-attribute preference weighting, "
        "budget boundary constraints, and natural-language AI insights into a unified residential decision-support platform."
    )
    p_intro2.paragraph_format.space_after = Pt(18)
    p_intro2.paragraph_format.line_spacing = 1.15

    # Section 2: Summary Table of All 15 Papers
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Comprehensive Literature Review Summary Table")
    r_h2.font.color.rgb = c_navy_rgb
    r_h2.font.bold = True

    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_names = ["Ref.", "Paper Title & Authors", "Journal / Publisher", "Year", "Methodology", "Category", "Status"]
    
    for idx, name in enumerate(hdr_names):
        hdr_cells[idx].text = name
        set_cell_background(hdr_cells[idx], c_navy_hex)
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[idx].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)
        if idx in [0, 3, 6]: p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for p_item in papers:
        row_cells = table.add_row().cells
        bg_hex = c_bg_hex if p_item["no"] % 2 == 1 else "FFFFFF"
        
        row_data = [
            (f"[{p_item['no']}]", True, WD_ALIGN_PARAGRAPH.CENTER),
            (f"{p_item['title']}\n({p_item['authors_short']})", True, WD_ALIGN_PARAGRAPH.LEFT),
            (p_item['journal'], False, WD_ALIGN_PARAGRAPH.LEFT),
            (str(p_item['year']), False, WD_ALIGN_PARAGRAPH.CENTER),
            (p_item['method_type'], False, WD_ALIGN_PARAGRAPH.LEFT),
            (p_item['category'], False, WD_ALIGN_PARAGRAPH.LEFT),
            (p_item['status'], True, WD_ALIGN_PARAGRAPH.CENTER)
        ]

        for idx, (text, is_bold, align) in enumerate(row_data):
            set_cell_background(row_cells[idx], bg_hex)
            set_cell_margins(row_cells[idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[idx].paragraphs[0]
            p.alignment = align
            run = p.add_run(text)
            run.font.size = Pt(8.5)
            run.font.bold = is_bold
            if idx == 6: run.font.color.rgb = c_emerald_rgb

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(18)

    # Section 3: Research Methodology Used in Each Paper
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Analysis of Methodologies Used in Literature")
    r_h3.font.color.rgb = c_navy_rgb

    p_method_overview = doc.add_paragraph(
        "Across the 15 surveyed papers, methodologies fall into seven distinct analytical paradigms:\n"
        "1. Discrete Choice Modeling & Econometrics (Schirmer et al., 2014; Walker & Li, 2007): Utilizes Multinomial Logit (MNL) and Latent Class models to estimate household utility utility coefficients from survey data.\n"
        "2. Stated Preference & Hierarchical Conjoint Analysis (Molin et al., 1996): Employs experimental trade-off scenarios to decompose consumer housing decisions into hierarchical utility trees.\n"
        "3. GIS Multi-Criteria Evaluation (Malczewski, 2006; Chandio et al., 2014): Integrates spatial overlay analysis in ArcGIS with weight matrices to produce continuous land suitability maps.\n"
        "4. Multi-Criteria Decision Making (MCDM) Algorithms (Zolfani et al., 2013; Pamučar et al., 2018): Employs Fuzzy AHP, WASPAS, and FUCOM optimization to resolve conflicting multi-attribute evaluation criteria.\n"
        "5. Content-Based Recommendation & Spatial Decay (Yuan et al., 2020; Wu et al., 2021): Combines Multi-Attribute Utility Theory (MAUT) with geographic distance decay functions to rank properties.\n"
        "6. Hedonic Price Modeling & Computer Vision (Zhang & Dong, 2018): Combines spatial hedonic regression with street-view image analysis to measure amenity impacts on housing valuations.\n"
        "7. Machine Learning & Explainable AI Frameworks (Baldominos et al., 2018; Dwivedi et al., 2021): Uses supervised regression trees (Random Forest) and Explainable AI (XAI) frameworks for automated evaluation."
    )
    p_method_overview.paragraph_format.line_spacing = 1.15
    p_method_overview.paragraph_format.space_after = Pt(18)

    # Section 4: What Each Paper Contributes & Limitations
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Detailed Review & Contributions of Individual Papers")
    r_h4.font.color.rgb = c_navy_rgb

    for p_item in papers:
        h_paper = doc.add_heading(level=2)
        r_hp = h_paper.add_run(f"[{p_item['no']}] {p_item['title']}")
        r_hp.font.color.rgb = c_blue_rgb

        p_det = doc.add_paragraph()
        p_det.paragraph_format.line_spacing = 1.15
        p_det.paragraph_format.space_after = Pt(12)
        
        p_det.add_run("• Authors & Year: ").bold = True
        p_det.add_run(f"{p_item['authors']} ({p_item['year']})\n")
        
        p_det.add_run("• Journal / Conference: ").bold = True
        p_det.add_run(f"{p_item['journal']} (DOI: {p_item['doi']})\n")
        
        p_det.add_run("• Research Problem: ").bold = True
        p_det.add_run(f"{p_item['proposes']}\n")
        
        p_det.add_run("• Methodology: ").bold = True
        p_det.add_run(f"{p_item['method']}\n")
        
        p_det.add_run("• Main Findings: ").bold = True
        p_det.add_run(f"{p_item['findings']}\n")
        
        p_det.add_run("• Limitation: ").bold = True
        p_det.add_run(f"{p_item['limitation']}\n")

        p_det.add_run("• Relevance to HomeWise AI: ").bold = True
        p_det.add_run(f"{p_item['relevance']}")

    # Section 5: Overall Research Gap
    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("5. Identification of Overall Research Gap")
    r_h5.font.color.rgb = c_navy_rgb

    p_gap = doc.add_paragraph(
        "Existing literature in urban economics, GIS spatial analysis, and multi-criteria decision-making (MCDM) has established "
        "that home-buying decisions are fundamentally multi-attribute spatial problems. Existing research addresses different parts of "
        "location and residential decision-making, such as GIS suitability analysis, MCDM ranking, environmental factors, infrastructure planning, "
        "machine-learning feature weighting, and consumer decision factors. However, these approaches are generally studied separately or for "
        "different applications.\n\n"
        "Specifically, three major research gaps persist in current systems:\n"
        "1. Domain Fragmentation: GIS suitability tools focus on municipal urban planning, MCDM studies focus on theoretical mathematical optimization, "
        "and real estate recommender systems focus on individual property listings rather than neighborhood liveability synthesis.\n"
        "2. Absence of Consumer-Facing Interactivity: Traditional spatial decision support relies on complex desktop GIS software (ArcGIS) "
        "requiring expert operators rather than lightweight, interactive web-map interfaces tailored for home buyers.\n"
        "3. Opacity of Black-Box Machine Learning: Modern AI real estate models output numerical predictions without natural-language "
        "explanations explaining *why* a specific location was recommended or highlighting critical due diligence trade-offs."
    )
    p_gap.paragraph_format.line_spacing = 1.15
    p_gap.paragraph_format.space_after = Pt(18)

    # Section 6: HomeWise AI Proposed Contribution
    h6 = doc.add_heading(level=1)
    r_h6 = h6.add_run("6. HomeWise AI Proposed Contribution")
    r_h6.font.color.rgb = c_navy_rgb

    p_contrib = doc.add_paragraph(
        "Existing studies generally address these components separately, creating an opportunity for an integrated personalized "
        "residential decision-support approach. HomeWise AI bridges these disconnected research silos by synthesizing six core elements "
        "into a single web application:\n\n"
        "• Buyer Preferences: Dynamic checkbox preference toggles allowing buyers to prioritize schools, hospitals, markets, safety, connectivity, air quality, flood safety, and future infrastructure.\n"
        "• Budget Constraints: Explicit min-max pricing boundary evaluation (in ₹ Crore) ensuring recommendations remain financially feasible.\n"
        "• Multiple Neighborhood Factors: Integrated evaluation across 9 synthetic liveability metrics per locality.\n"
        "• Geospatial Web Visualization: Interactive client-side Leaflet GIS mapping with 40 Mumbai Metropolitan Region locality markers.\n"
        "• Weighted Locality Ranking: Transparent preference-weighted recommendation scoring algorithm yielding ranked Top 5 matches.\n"
        "• AI-Based Explanation: Dynamic synthesis of natural-language intelligence ('Prototype AI Insight') providing contextual recommendations, due-diligence caveats, and verdict summaries."
    )
    p_contrib.paragraph_format.line_spacing = 1.15
    p_contrib.paragraph_format.space_after = Pt(24)

    # Section 7: References
    h7 = doc.add_heading(level=1)
    r_h7 = h7.add_run("7. Academic References")
    r_h7.font.color.rgb = c_navy_rgb

    for p_item in papers:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.left_indent = Inches(0.4)
        p_ref.paragraph_format.first_line_indent = Inches(-0.4)
        p_ref.paragraph_format.space_after = Pt(6)
        
        r_num = p_ref.add_run(f"[{p_item['no']}] ")
        r_num.bold = True
        
        ref_str = f"{p_item['authors']}, \"{p_item['title']}\", {p_item['journal']}, {p_item['year']}."
        p_ref.add_run(ref_str + " ")
        
        r_doi = p_ref.add_run(f"DOI: {p_item['doi']}")
        r_doi.font.color.rgb = c_blue_rgb

    doc_path = "HomeWise_AI_15_Paper_Literature_Review.docx"
    doc.save(doc_path)
    print(f"Document saved successfully to {doc_path}")

if __name__ == "__main__":
    create_document()
